# -*- coding: utf-8 -*-
"""Outreach Platform — app di invio email massivo (schema `mass`).

App separata che usa lo stesso database di `cv-scouting-lusha` ma in uno schema
dedicato: NON tocca le tabelle dell'app esistente.
"""
import io
import re
import secrets
import unicodedata
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import quote_plus

import pandas as pd
from fastapi import FastAPI, Request, Form, File, UploadFile, Depends
from fastapi.responses import HTMLResponse, RedirectResponse, FileResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

from app import ai, learning, mailer, models, replies, scheduler
from app.config import (APP_HOST, APP_PORT, AUTOSEND, DAILY_LIMIT, SCHEDULER_TOKEN,
                        SEND_WINDOW_END, SEND_WINDOW_START)
from app.db import SessionLocal, init_db
from app.email_builder import (
    LINK, LOGO, SUBJECT, build_apertura, build_email, build_email_html, build_saluto,
    dividi_nome, get_email_template, normalizza_testo, pulisci_nome, saluto_orario,
)
from app.security import COOKIE, check_password, make_token, read_token

BASE = Path(__file__).resolve().parent
STATIC = BASE.parent / "static"
TEMPLATES = BASE.parent / "templates"

ALIASES = {
    "Azienda": ["azienda", "nome azienda", "ragione sociale", "denominazione",
                "denominazione aziendale", "societa", "impresa", "ditta", "cliente",
                "organizzazione", "company", "company name", "organization", "account"],
    "Contatto": ["contatto", "nome", "nome e cognome", "nome completo", "nome referente",
                 "contact", "contact name", "full name", "referente", "contatto hr",
                 "contatto aziendale", "persona", "destinatario"],
    "Cognome": ["cognome", "cognomi", "last name", "surname"],
    "Email": ["email", "e mail", "indirizzo email", "indirizzo mail", "indirizzo e mail",
              "mail", "email aziendale", "mail aziendale", "email contatto", "contatto email",
              "email address", "work email"],
    "Timing": ["timing", "timing outreach", "tempistica", "quando", "finestra temporale"],
    "Flag": ["flag", "alert", "correzione", "correzione alert", "avviso", "attenzione", "verifica"],
    "Contesto": ["contesto", "settore", "settore prodotto", "settore merceologico",
                 "settore di appartenenza", "ambito", "comparto", "mercato di riferimento",
                 "industry", "sector"],
    "Gancio": ["gancio", "hook", "aggancio", "appiglio", "spunto", "spunto per l'email",
               "spunto email", "spunto per email", "idea email", "riferimento",
               "riferimento mercato"],
    "RicercaAttivita": ["attivita", "attivita in italia", "attivita italia", "attivita azienda",
                        "activity", "business", "business in italia", "business italia",
                        "business in italy", "descrizione business", "business description",
                        "cosa fa", "di cosa si occupa", "contesto italia", "contesto e mail",
                        "contesto email", "gancio contesto italia"],
    "RicercaCompetenze": ["competenze", "competenze da mappare", "competenze chiave",
                          "competenze ricercate", "competenze target", "competenze tecniche",
                          "professionalita", "profili", "profili da mappare", "ruoli da mappare",
                          "skills", "skill", "competencies", "profiles"],
    "RicercaSegnale": ["segnale", "segnale recente", "segnale recente perche ora",
                       "segnale di attivita", "segnale perche ora", "evento recente",
                       "eventi recenti", "novita", "attualita", "notizie", "news", "signal",
                       "recent signal", "trigger", "trigger event"],
    "AlertText": ["alert", "testo alert", "alert text", "correzione", "correzione alert",
                  "nota alert", "nota correzione", "segnalazione"],
}

from contextlib import asynccontextmanager


def _scheduler_loop():
    """Avvia automaticamente gli invii programmati, dentro la finestra oraria e sotto il tetto."""
    import threading
    import time as _time
    while True:
        try:
            _time.sleep(45)
            if not AUTOSEND:
                continue
            if mailer.STATE.get("running") or not mailer.nella_finestra() or mailer.tetto_raggiunto():
                continue
            db = SessionLocal()
            try:
                for camp in db.query(models.Campaign).filter_by(auto_send=True).all():
                    rows = _ready_rows(db, camp)
                    if not rows:
                        camp.auto_send = False
                        db.commit()
                        continue
                    with mailer.LOCK:
                        if mailer.STATE.get("running"):
                            break
                    mailer.STOP.clear()
                    threading.Thread(target=mailer.run_send,
                                     args=(rows, camp.channel, True, 1, True),
                                     daemon=True).start()
                    break
            finally:
                db.close()
        except Exception:
            pass


@asynccontextmanager
async def lifespan(_app):
    init_db()
    import threading
    threading.Thread(target=_scheduler_loop, daemon=True).start()
    yield


app = FastAPI(title="Schoenbech Outreach Platform", lifespan=lifespan)
app.mount("/static", StaticFiles(directory=str(STATIC)), name="static")
templates = Jinja2Templates(directory=str(TEMPLATES))


# ---------------------------------------------------------------------------
# Helper import
# ---------------------------------------------------------------------------
def _norm(s):
    s = str(s).lower().strip()
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def normalize_flag(v):
    n = _norm(v)
    if not n:
        return "OK"
    if "nessuna correz" in n or "nessun alert" in n:
        return "OK"
    if "sensibil" in n or "delicat" in n:
        return "Sensibile"
    if "verific" in n or "attenzione" in n:
        return "Verifica"
    if "correz" in n or "troppo" in n or "incomplet" in n or "generico" in n or "stretto" in n:
        return "Correzione"
    return "OK"


def _match_columns(header):
    """Mappa le colonne del file sui campi attesi.

    Prima cerca i nomi esatti (come da ALIASES); per i campi ancora vuoti
    accetta anche nomi che contengono una variante (es. "Settore / Prodotto"
    -> Contesto, "Business in Italia" -> RicercaAttivita).
    """
    norm = [_norm(c) for c in header]
    wanted = {t: ({_norm(t)} | {_norm(a) for a in al}) for t, al in ALIASES.items()}
    colmap = {}
    for target, varianti in wanted.items():
        for j, nc in enumerate(norm):
            if nc in varianti:
                colmap[target] = j
                break
    for target, varianti in wanted.items():
        if target in colmap:
            continue
        for j, nc in enumerate(norm):
            if nc and any(v and v in nc and not (v == "nome" and "cognome" in nc)
                          for v in varianti):
                colmap[target] = j
                break
    return colmap


def parse_uploaded(content, filename):
    bio = io.BytesIO(content)
    if (filename or "").lower().endswith(".csv"):
        sheets = {"csv": pd.read_csv(bio, header=None, dtype=str, sep=None,
                                     engine="python", keep_default_na=False)}
    else:
        sheets = pd.read_excel(bio, header=None, dtype=str, keep_default_na=False, sheet_name=None)
    best, best_score = [], -1
    for _n, raw in sheets.items():
        rows, score = _parse_sheet(raw.fillna(""))
        if score > best_score:
            best, best_score = rows, score
    return best


def _parse_sheet(raw):
    if len(raw) == 0:
        return [], 0
    alias_norm = set()
    for k, vs in ALIASES.items():
        alias_norm.add(_norm(k))
        for v in vs:
            alias_norm.add(_norm(v))
    hidx = 0
    for r in range(min(12, len(raw))):
        vals = [_norm(v) for v in raw.iloc[r].tolist()]
        if sum(1 for v in vals if v in alias_norm) >= 2 or "azienda" in vals or "email" in vals:
            hidx = r
            break
    header = [str(x).strip() for x in raw.iloc[hidx].tolist()]
    data = raw.iloc[hidx + 1:].reset_index(drop=True)
    colmap = _match_columns(header)
    rows = []
    for pos in range(len(data)):
        vals = data.iloc[pos].tolist()

        def g(t):
            j = colmap.get(t)
            return str(vals[j]).strip() if j is not None else ""

        az, em = g("Azienda"), g("Email")
        if not az and not em:
            continue
        contatto = g("Contatto")
        if g("Cognome") and colmap.get("Cognome") != colmap.get("Contatto"):
            contatto = f"{contatto} {g('Cognome')}".strip()
        rows.append({
            "azienda": az, "contatto": contatto, "email": em,
            "timing": g("Timing"), "flag": normalize_flag(g("Flag")),
            "contesto": g("Contesto"), "gancio": g("Gancio"),
            "ricerca_attivita": g("RicercaAttivita"), "ricerca_competenze": g("RicercaCompetenze"),
            "ricerca_segnale": g("RicercaSegnale"), "alert_text": g("AlertText"),
        })
    return rows, len(colmap) + (3 if "Email" in colmap else 0)


# ---------------------------------------------------------------------------
# Auth
# ---------------------------------------------------------------------------
class _Redirect(Exception):
    def __init__(self, url):
        self.url = url


@app.exception_handler(_Redirect)
def _redirect_handler(request, exc):
    return RedirectResponse(exc.url, status_code=303)


def require_user(request: Request):
    user = read_token(request.cookies.get(COOKIE))
    if not user:
        raise _Redirect("/login")
    return user


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    return templates.TemplateResponse(request, "login.html", {"error": None})


@app.post("/login")
def login_submit(password: str = Form("")):
    if check_password(password):
        resp = RedirectResponse("/", status_code=303)
        resp.set_cookie(COOKIE, make_token(), httponly=True, samesite="lax", max_age=60 * 60 * 24 * 30)
        return resp
    return RedirectResponse("/login?e=1", status_code=303)


@app.get("/logout")
def logout():
    resp = RedirectResponse("/login", status_code=303)
    resp.delete_cookie(COOKIE)
    return resp


@app.get("/logo.png")
def logo_png():
    return FileResponse(str(LOGO))


# ---------------------------------------------------------------------------
# Campagne
# ---------------------------------------------------------------------------
@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request, user=Depends(require_user)):
    db = SessionLocal()
    try:
        camps = db.query(models.Campaign).order_by(models.Campaign.created_at.desc()).all()
        data = []
        for c in camps:
            n = db.query(models.CampaignCompany).filter_by(campaign_id=c.campaign_id).count()
            data.append({"c": c, "n": n})
        risposte = replies.ultime_risposte(db)
    finally:
        db.close()
    return templates.TemplateResponse(request, "index.html", {
        "campaigns": data, "risposte": risposte, "scan_state": replies.STATE, "user": user})


@app.post("/campaigns/new")
def create_campaign(name: str = Form(""), user=Depends(require_user)):
    db = SessionLocal()
    try:
        c = models.Campaign(name=name.strip() or "Nuova campagna", created_by=user.get("u"))
        db.add(c)
        db.commit()
        cid = c.campaign_id
    finally:
        db.close()
    return RedirectResponse(f"/campaign/{cid}", status_code=303)


@app.get("/campaign/{cid}", response_class=HTMLResponse)
def campaign_detail(request: Request, cid: str, user=Depends(require_user)):
    db = SessionLocal()
    try:
        camp = db.get(models.Campaign, cid)
        if not camp:
            return RedirectResponse("/", status_code=303)
        rows = []
        for cc in db.query(models.CampaignCompany).filter_by(campaign_id=camp.campaign_id).all():
            comp = db.get(models.Company, cc.company_id)
            con = db.get(models.Contact, cc.contact_id) if cc.contact_id else None
            rows.append({
                "cc_id": str(cc.campaign_company_id),
                "azienda": comp.legal_name if comp else "",
                "saluto": cc.saluto or "",
                "email": (con.email if con and con.email else ""),
                "timing": cc.timing or "", "flag": cc.flag or "OK",
                "tipo_gancio": cc.tipo_gancio or "", "gancio": cc.gancio or "",
                "stato": cc.draft_status or "Bozza",
                "motivo": cc.escluso_motivo or "",
                "risposto": con.risposto_at if con else None,
                "risposta_oggetto": (con.risposta_oggetto if con else "") or "",
            })
        counts = {}
        for r in rows:
            counts[r["stato"]] = counts.get(r["stato"], 0) + 1
        ready = _ready_rows(db, camp)
        due = _recall_rows(db, camp)
        n_ricerca = sum(1 for cc in db.query(models.CampaignCompany)
                        .filter_by(campaign_id=camp.campaign_id, draft_status="Bozza").all()
                        if (cc.ricerca_attivita or cc.ricerca_competenze
                            or cc.ricerca_segnale or cc.alert_text
                            or (cc.contesto or "").strip() or (cc.gancio or "").strip()))
        n_correzioni = learning.conta_correzioni(db)
        n_risposte = replies.conta_risposte(db, camp.campaign_id)
    finally:
        db.close()
    return templates.TemplateResponse(request, "campaign.html", {
        "camp": camp, "rows": rows, "counts": counts, "tot": len(rows),
        "n_ready": len(ready), "n_due": len(due), "n_ricerca": n_ricerca,
        "n_correzioni": n_correzioni, "n_risposte": n_risposte,
        "ai_state": ai.AI_STATE, "user": user,
    })


@app.post("/import/{cid}")
def import_file(cid: str, file: UploadFile = File(...), user=Depends(require_user)):
    content = file.file.read()
    try:
        parsed = parse_uploaded(content, file.filename or "")
    except Exception as e:
        return RedirectResponse(f"/campaign/{cid}?msg=Errore+import:+{e}", status_code=303)
    db = SessionLocal()
    added = updated = duplicati = 0
    try:
        camp = db.get(models.Campaign, cid)
        if not camp:
            return RedirectResponse("/", status_code=303)
        for r in parsed:
            norm = _norm(r["azienda"])
            comp = db.query(models.Company).filter_by(normalized_name=norm).first()
            if not comp:
                comp = models.Company(legal_name=r["azienda"] or "(senza nome)",
                                      normalized_name=norm, source="import")
                db.add(comp)
                db.flush()
            con = None
            if r["email"]:
                con = (db.query(models.Contact)
                       .filter_by(company_id=comp.company_id,
                                  email_normalized=r["email"].lower()).first())
                if not con:
                    nomi, cognome = dividi_nome(r["contatto"])
                    con = models.Contact(company_id=comp.company_id, email=r["email"],
                                         email_normalized=r["email"].lower(),
                                         first_name=nomi, last_name=cognome, source="import")
                    db.add(con)
                    db.flush()
            # riga campagna = stessa azienda + stesso interlocutore (più contatti per azienda ammessi)
            cc = (db.query(models.CampaignCompany)
                  .filter_by(campaign_id=camp.campaign_id, company_id=comp.company_id)
                  .filter(models.CampaignCompany.contact_id == (con.contact_id if con else None))
                  .first())
            if not cc and con:  # adotta una riga esistente senza contatto
                cc = (db.query(models.CampaignCompany)
                      .filter_by(campaign_id=camp.campaign_id, company_id=comp.company_id)
                      .filter(models.CampaignCompany.contact_id.is_(None)).first())
            nuovo = False
            if not cc:
                cc = models.CampaignCompany(campaign_id=camp.campaign_id, company_id=comp.company_id)
                db.add(cc)
                db.flush()
                nuovo = True
            cc.contact_id = con.contact_id if con else None
            if r["email"] and con:
                con.email = r["email"]
                con.email_normalized = r["email"].lower()
            if r["contesto"] and not (cc.contesto or "").strip():
                cc.contesto = r["contesto"]
            if r["gancio"] and not (cc.gancio or "").strip():
                cc.gancio = r["gancio"]
            if r["timing"]:
                cc.timing = r["timing"]
            if r["flag"]:
                cc.flag = r["flag"]
            if r["contatto"]:
                cc.saluto = build_saluto(r["contatto"])
                if con:
                    nomi, cognome = dividi_nome(r["contatto"])
                    con.first_name, con.last_name = nomi, cognome
            if r.get("ricerca_attivita"):
                cc.ricerca_attivita = r["ricerca_attivita"]
            if r.get("ricerca_competenze"):
                cc.ricerca_competenze = r["ricerca_competenze"]
            if r.get("ricerca_segnale"):
                cc.ricerca_segnale = r["ricerca_segnale"]
            if r.get("alert_text"):
                cc.alert_text = r["alert_text"]
            cc.paragrafo_apertura = build_apertura(cc.contesto or "", cc.gancio or "")
            em = (con.email if con else "") or ""
            # Un'esclusione (manuale o per duplicato) non si riapre da sola.
            if cc.draft_status != "Esclusa":
                cc.draft_status = "Pronta" if (cc.contesto and cc.gancio and em) else "Bozza"
                # Memoria degli invii: se questo indirizzo e' gia' stato raggiunto in
                # un'altra campagna, la riga entra ma resta esclusa con il motivo.
                precedente = learning.gia_contattato(db, em, camp.campaign_id)
                if precedente:
                    quando, campagna = precedente
                    data_txt = quando.strftime("%d/%m/%Y") if quando else "data sconosciuta"
                    cc.draft_status = "Esclusa"
                    cc.escluso_motivo = (f"Gia' contattato il {data_txt}"
                                         + (f" (campagna \u00ab{campagna}\u00bb)" if campagna else ""))
                    duplicati += 1
            if nuovo:
                added += 1
            else:
                updated += 1
        db.commit()
    finally:
        db.close()
    msg = f"Aggiunti+{added},+aggiornati+{updated}"
    if duplicati:
        msg += f",+esclusi+{duplicati}+gia+contattati"
    return RedirectResponse(f"/campaign/{cid}?msg={msg}", status_code=303)


@app.post("/campaign/{cid}/rename")
def rename_campaign(cid: str, name: str = Form(""), back: str = Form("/"),
                    user=Depends(require_user)):
    db = SessionLocal()
    try:
        camp = db.get(models.Campaign, cid)
        if camp and name.strip():
            camp.name = name.strip()
            db.commit()
    finally:
        db.close()
    return RedirectResponse(back or f"/campaign/{cid}", status_code=303)


@app.get("/campaign/{cid}/report.xlsx")
def campaign_report(cid: str, user=Depends(require_user)):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    db = SessionLocal()
    try:
        camp = db.get(models.Campaign, cid)
        if not camp:
            return RedirectResponse("/", status_code=303)
        wb = Workbook()
        ws = wb.active
        ws.title = "Invii"
        ws.append(["Azienda", "Referente", "Email", "Ruolo", "Stato bozza",
                   "Email 1", "Data/ora invio 1", "Canale 1", "Errore 1",
                   "Recall", "Data/ora recall", "Errore recall",
                   "Ha risposto", "Data risposta", "Oggetto risposta"])

        def esito(m):
            if not m:
                return "Non inviata"
            return "Inviata" if m.status == "sent" else "Errore"

        def quando(m):
            return m.sent_at.strftime("%d/%m/%Y %H:%M") if (m and m.sent_at) else ""

        def errore(m, cc):
            if not m or m.status == "sent":
                return ""
            ev = (db.query(models.OutreachEvent)
                  .filter(models.OutreachEvent.campaign_id == cc.campaign_id,
                          models.OutreachEvent.contact_id == cc.contact_id,
                          models.OutreachEvent.event_type == "failed")
                  .order_by(models.OutreachEvent.occurred_at.desc()).first())
            return str((ev.payload or {}).get("error", ""))[:200] if ev else ""

        for cc in db.query(models.CampaignCompany).filter_by(campaign_id=camp.campaign_id).all():
            comp = db.get(models.Company, cc.company_id)
            con = db.get(models.Contact, cc.contact_id) if cc.contact_id else None
            m1 = (db.query(models.OutreachMessage)
                  .filter_by(campaign_company_id=cc.campaign_company_id, message_no=1)
                  .order_by(models.OutreachMessage.created_at.desc()).first())
            m2 = (db.query(models.OutreachMessage)
                  .filter_by(campaign_company_id=cc.campaign_company_id, message_no=2)
                  .order_by(models.OutreachMessage.created_at.desc()).first())
            ws.append([comp.legal_name if comp else "", cc.saluto or "",
                       (con.email if con else ""), ((con.role if con else "") or ""),
                       cc.draft_status, esito(m1), quando(m1), (m1.channel if m1 else ""),
                       errore(m1, cc), esito(m2), quando(m2), errore(m2, cc),
                       ("Si" if (con and con.risposto_at) else ""),
                       (con.risposto_at.strftime("%d/%m/%Y %H:%M")
                        if (con and con.risposto_at) else ""),
                       ((con.risposta_oggetto if con else "") or "")])
        for i, w in enumerate([28, 26, 30, 24, 12, 12, 18, 10, 42, 12, 18, 42,
                               12, 18, 46], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        bio = io.BytesIO()
        wb.save(bio)
        nome = "report_" + re.sub(r"[^A-Za-z0-9_-]+", "_", camp.name)[:40] + ".xlsx"
        return Response(content=bio.getvalue(),
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="{nome}"'})
    finally:
        db.close()


@app.post("/campaign/{cid}/delete")
def delete_campaign(cid: str, user=Depends(require_user)):
    db = SessionLocal()
    try:
        camp = db.get(models.Campaign, cid)
        if camp:
            ccs = db.query(models.CampaignCompany).filter_by(campaign_id=camp.campaign_id).all()
            comp_ids = {cc.company_id for cc in ccs}
            db.query(models.CampaignCompany).filter_by(campaign_id=camp.campaign_id).delete()
            db.delete(camp)
            db.flush()
            for comp_id in comp_ids:
                used = db.query(models.CampaignCompany).filter_by(company_id=comp_id).count()
                if not used:
                    db.query(models.Contact).filter_by(company_id=comp_id).delete()
                    db.query(models.Company).filter_by(company_id=comp_id).delete()
            db.commit()
    finally:
        db.close()
    return RedirectResponse("/", status_code=303)


@app.post("/campaign/{cid}/fix-saluti")
def fix_saluti(cid: str, user=Depends(require_user)):
    db = SessionLocal()
    fixed = 0
    try:
        camp = db.get(models.Campaign, cid)
        if not camp:
            return RedirectResponse("/", status_code=303)
        for cc in db.query(models.CampaignCompany).filter_by(campaign_id=camp.campaign_id).all():
            con = db.get(models.Contact, cc.contact_id) if cc.contact_id else None
            if not con:
                continue
            fn = pulisci_nome(con.first_name or "")
            ln = pulisci_nome(con.last_name or "")
            if fn != (con.first_name or ""):
                con.first_name = fn
            if ln != (con.last_name or ""):
                con.last_name = ln
            nuovo = build_saluto(f"{fn} {ln}".strip())
            if nuovo and cc.saluto != nuovo:
                cc.saluto = nuovo
                fixed += 1
        db.commit()
    finally:
        db.close()
    return RedirectResponse(f"/campaign/{cid}?msg=Saluti+corretti:+{fixed}", status_code=303)


@app.post("/campaign/{cid}/genera-ai")
def genera_ai(cid: str, user=Depends(require_user)):
    if ai.AI_STATE.get("running"):
        return RedirectResponse(f"/campaign/{cid}?msg=Generazione+AI+gia+in+corso", status_code=303)
    db = SessionLocal()
    try:
        camp = db.get(models.Campaign, cid)
        if not camp:
            return RedirectResponse("/", status_code=303)
        rows = []
        for cc in db.query(models.CampaignCompany).filter_by(campaign_id=camp.campaign_id).all():
            if cc.draft_status != "Bozza":  # le righe gia' Pronte/Escluse non si rigenerano
                continue
            ha_ricerca = (cc.ricerca_attivita or cc.ricerca_competenze
                          or cc.ricerca_segnale or cc.alert_text)
            ha_attuale = (cc.contesto or "").strip() or (cc.gancio or "").strip()
            if not ha_ricerca and not ha_attuale:
                continue
            comp = db.get(models.Company, cc.company_id)
            nome = comp.legal_name if comp else ""
            data = {"azienda": nome, "attivita": cc.ricerca_attivita or "",
                    "competenze": cc.ricerca_competenze or "",
                    "segnale": cc.ricerca_segnale or "",
                    "alert": cc.alert_text or "", "timing": cc.timing or ""}
            if ha_attuale:
                data["versione_attuale"] = {"settore": (cc.contesto or "").strip(),
                                            "profili": (cc.gancio or "").strip()}
            rows.append({"cc_id": cc.campaign_company_id, "azienda": nome, "data": data,
                         "cutoff": cc.created_at})
    finally:
        db.close()
    if not rows:
        return RedirectResponse(f"/campaign/{cid}?msg=Nessun+dato+di+ricerca+da+elaborare", status_code=303)
    ai.AI_STOP.clear()
    import threading
    threading.Thread(target=ai.run_generate, args=(rows,), daemon=True).start()
    return RedirectResponse(f"/campaign/{cid}", status_code=303)


@app.get("/ai/status")
def ai_status(user=Depends(require_user)):
    with ai.AI_LOCK:
        return JSONResponse(dict(ai.AI_STATE))


@app.post("/ai/regen/{cc_id}")
def ai_regen(cc_id: str, back: str = Form("/"), user=Depends(require_user)):
    msg = ai.generate_one(cc_id)
    sep = "&" if "?" in back else "?"
    from urllib.parse import quote
    return RedirectResponse(f"{back}{sep}msg={quote(msg)}", status_code=303)


@app.get("/preview/{cc_id}", response_class=HTMLResponse)
def preview(request: Request, cc_id: str, user=Depends(require_user)):
    db = SessionLocal()
    try:
        cc = db.get(models.CampaignCompany, cc_id)
        if not cc:
            return RedirectResponse("/", status_code=303)
        comp = db.get(models.Company, cc.company_id)
        con = db.get(models.Contact, cc.contact_id) if cc.contact_id else None
        row = {"azienda": comp.legal_name if comp else "", "saluto": cc.saluto or "",
               "email": con.email if con else "", "contesto": cc.contesto or "",
               "gancio": cc.gancio or "", "flag": cc.flag or "OK", "stato": cc.draft_status,
               "paragrafo_apertura": cc.paragrafo_apertura or build_apertura(cc.contesto or "", cc.gancio or "")}
        tpl = get_email_template()
        subject, body = build_email(row, 1, tpl)
        cid = str(cc.campaign_id)
    finally:
        db.close()
    return templates.TemplateResponse(request, "preview.html",
                                      {"row": row, "subject": subject, "body": body, "i": cc_id, "cid": cid})


@app.get("/preview-html/{cc_id}", response_class=HTMLResponse)
def preview_html(cc_id: str, user=Depends(require_user)):
    db = SessionLocal()
    try:
        cc = db.get(models.CampaignCompany, cc_id)
        if not cc:
            return HTMLResponse("<p>Non trovato</p>")
        row = {"saluto": cc.saluto or "", "contesto": cc.contesto or "", "gancio": cc.gancio or ""}
    finally:
        db.close()
    tpl = get_email_template()
    return HTMLResponse(build_email_html(row, logo="/logo.png", message_no=1, tpl=tpl))


@app.post("/stato/{cc_id}")
def set_stato(cc_id: str, stato: str = Form("Bozza"), back: str = Form("/"), user=Depends(require_user)):
    db = SessionLocal()
    try:
        cc = db.get(models.CampaignCompany, cc_id)
        if cc and stato in ("Bozza", "Pronta", "Esclusa"):
            cc.draft_status = stato
            db.commit()
    finally:
        db.close()
    return RedirectResponse(back, status_code=303)


@app.post("/set-email/{cc_id}")
def set_email(cc_id: str, email: str = Form(""), back: str = Form("/"), user=Depends(require_user)):
    db = SessionLocal()
    try:
        cc = db.get(models.CampaignCompany, cc_id)
        if cc and cc.contact_id:
            con = db.get(models.Contact, cc.contact_id)
            if con:
                con.email = email.strip()
                con.email_normalized = email.strip().lower()
                if cc.contesto and cc.gancio and email.strip() and cc.draft_status == "Bozza":
                    cc.draft_status = "Pronta"
                db.commit()
    finally:
        db.close()
    return RedirectResponse(back, status_code=303)


@app.post("/apertura/{cc_id}")
def set_apertura(cc_id: str, saluto: str = Form(""), contesto: str = Form(""), gancio: str = Form(""),
                 condividi: str = Form(""), back: str = Form("/"), user=Depends(require_user)):
    db = SessionLocal()
    try:
        cc = db.get(models.CampaignCompany, cc_id)
        if cc:
            # Prima di sovrascrivere: se l'operatore ha cambiato la proposta AI,
            # la differenza viene salvata e insegnera' alle generazioni future.
            comp = db.get(models.Company, cc.company_id)
            learning.registra_correzione(db, cc, contesto, gancio,
                                         azienda=comp.legal_name if comp else "")
            cc.saluto = saluto.strip()
            cc.contesto = contesto.strip()
            cc.gancio = gancio.strip()
            cc.paragrafo_apertura = build_apertura(cc.contesto or "", cc.gancio or "")
            con = db.get(models.Contact, cc.contact_id) if cc.contact_id else None
            em = (con.email if con else "") or ""
            cc.draft_status = "Pronta" if (cc.contesto and cc.gancio and em) else "Bozza"
            if condividi == "on":
                fratelli = (db.query(models.CampaignCompany)
                            .filter_by(campaign_id=cc.campaign_id, company_id=cc.company_id)
                            .filter(models.CampaignCompany.campaign_company_id != cc.campaign_company_id)
                            .all())
                for altro in fratelli:
                    altro.contesto = cc.contesto
                    altro.gancio = cc.gancio
                    altro.paragrafo_apertura = cc.paragrafo_apertura
                    con2 = db.get(models.Contact, altro.contact_id) if altro.contact_id else None
                    em2 = (con2.email if con2 else "") or ""
                    altro.draft_status = "Pronta" if (altro.contesto and altro.gancio and em2) else "Bozza"
            db.commit()
    finally:
        db.close()
    return RedirectResponse(back or "/", status_code=303)


@app.get("/testo", response_class=HTMLResponse)
def testo_page(request: Request, user=Depends(require_user)):
    tpl = get_email_template()
    return templates.TemplateResponse(request, "testo.html", {"tpl": tpl, "link": LINK})


@app.post("/testo")
def testo_save(subject: str = Form(""), corpo: str = Form(""), recall_text: str = Form(""),
               user=Depends(require_user)):
    db = SessionLocal()
    try:
        row = db.get(models.EmailTemplate, 1)
        if not row:
            row = models.EmailTemplate(template_id=1)
            db.add(row)
        row.subject = subject.strip() or None
        # Un indirizzo a scadenza incollato qui verrebbe spedito per settimane:
        # diventa subito il segnaposto, cosi' resta una sola fonte per il link.
        row.corpo = normalizza_testo(corpo.strip()) or None
        row.recall_text = normalizza_testo(recall_text.strip()) or None
        db.commit()
        cambiato = (row.corpo != (corpo.strip() or None)
                    or row.recall_text != (recall_text.strip() or None))
    finally:
        db.close()
    msg = ("Testo+salvato+(link+a+scadenza+sostituito+con+il+segnaposto)"
           if cambiato else "Testo+salvato")
    return RedirectResponse(f"/testo?msg={msg}", status_code=303)


@app.post("/testo/reset")
def testo_reset(user=Depends(require_user)):
    db = SessionLocal()
    try:
        row = db.get(models.EmailTemplate, 1)
        if row:
            db.delete(row)
            db.commit()
    finally:
        db.close()
    return RedirectResponse("/testo?msg=Testo+ripristinato", status_code=303)


# ---------------------------------------------------------------------------
# Email singola (inserimento manuale di un referente)
# ---------------------------------------------------------------------------
@app.get("/singola", response_class=HTMLResponse)
def singola_page(request: Request, user=Depends(require_user)):
    return templates.TemplateResponse(request, "singola.html", {})


@app.post("/singola/crea")
def singola_crea(titolo: str = Form(""), nome: str = Form(""), cognome: str = Form(""),
                 azienda: str = Form(""), email: str = Form(""), ruolo: str = Form(""),
                 info: str = Form(""), competenze: str = Form(""),
                 settore: str = Form(""), profili: str = Form(""),
                 user=Depends(require_user)):
    nome, cognome = nome.strip(), cognome.strip()
    azienda, email = azienda.strip(), email.strip()
    if not (nome and cognome and azienda and email):
        return RedirectResponse("/singola?msg=Compila+nome,+cognome,+azienda+ed+email", status_code=303)
    db = SessionLocal()
    ai_msg, ccid = "", None
    try:
        camp = db.query(models.Campaign).filter_by(name="Singoli").first()
        if not camp:
            camp = models.Campaign(name="Singoli", created_by=user.get("u"))
            db.add(camp)
            db.flush()
        norm = _norm(azienda)
        comp = db.query(models.Company).filter_by(normalized_name=norm).first()
        if not comp:
            comp = models.Company(legal_name=azienda, normalized_name=norm, source="singola")
            db.add(comp)
            db.flush()
        con = (db.query(models.Contact)
               .filter_by(company_id=comp.company_id, email_normalized=email.lower()).first())
        if not con:
            con = models.Contact(company_id=comp.company_id, email=email,
                                 email_normalized=email.lower(), first_name=nome, last_name=cognome,
                                 role=ruolo.strip(), source="singola")
            db.add(con)
            db.flush()
        else:
            con.email, con.first_name, con.last_name = email, nome, cognome
            if ruolo.strip():
                con.role = ruolo.strip()
        contesto, gancio = settore.strip(), profili.strip()
        ai_ctx = ai_gan = ""
        if not (contesto and gancio) and (info.strip() or competenze.strip()):
            try:
                esempi = learning.blocco_prompt(
                    learning.esempi_validi(learning.carica_correzioni(db),
                                           datetime.now(timezone.utc)))
                out = ai.generate_personalization({
                    "azienda": azienda, "attivita": info.strip(), "competenze": competenze.strip(),
                    "segnale": "", "alert": "", "timing": ""}, esempi)
                ai_ctx = (out.get("contesto") or "").strip()
                ai_gan = (out.get("gancio") or "").strip()
                contesto = ai_ctx or contesto
                gancio = ai_gan or gancio
            except Exception as e:
                ai_msg = f" (AI non disponibile: {e})"
        cc = (db.query(models.CampaignCompany)
              .filter_by(campaign_id=camp.campaign_id, company_id=comp.company_id).first())
        if not cc:
            cc = models.CampaignCompany(campaign_id=camp.campaign_id, company_id=comp.company_id)
            db.add(cc)
        cc.contact_id = con.contact_id
        if contesto:
            cc.contesto = contesto
        if gancio:
            cc.gancio = gancio
        if ai_ctx or ai_gan:  # tiene la proposta AI per riconoscere le correzioni
            cc.contesto_ai, cc.gancio_ai = ai_ctx, ai_gan
        cc.saluto = build_saluto(f"{nome} {cognome}", titolo=titolo.strip())
        cc.paragrafo_apertura = build_apertura(cc.contesto or "", cc.gancio or "")
        cc.draft_status = "Pronta" if (cc.contesto and cc.gancio and email) else "Bozza"
        db.commit()
        ccid = cc.campaign_company_id
    except Exception as e:
        ai_msg = f"Errore: {e}"
    finally:
        db.close()
    if not ccid:
        return RedirectResponse(f"/singola?msg={quote_plus(ai_msg or 'Errore nella creazione')}", status_code=303)
    return RedirectResponse(f"/preview/{ccid}?msg={quote_plus('Email creata' + ai_msg)}", status_code=303)


@app.post("/singola/invia/{cc_id}")
def singola_invia(cc_id: str, channel: str = Form("Gmail"), salva_copia: str = Form(""),
                  user=Depends(require_user)):
    db = SessionLocal()
    msg = ""
    try:
        cc = db.get(models.CampaignCompany, cc_id)
        if not cc:
            return RedirectResponse("/", status_code=303)
        camp = db.get(models.Campaign, cc.campaign_id)
        row = _row_from_cc(db, cc, camp)
        gia = (db.query(models.OutreachMessage)
               .filter_by(campaign_company_id=cc.campaign_company_id, message_no=1, status="sent").first())
        if gia:
            msg = "Questa email risulta gia inviata"
        elif not row["email"]:
            msg = "Manca l'indirizzo email"
        else:
            msg = mailer.send_singolo(row, channel, salva_copia=(salva_copia == "on"))
    finally:
        db.close()
    return RedirectResponse(f"/preview/{cc_id}?msg={quote_plus(msg)}", status_code=303)


# ---------------------------------------------------------------------------
# Invio / Recall — selezione righe
# ---------------------------------------------------------------------------
def _row_from_cc(db, cc, camp):
    comp = db.get(models.Company, cc.company_id)
    con = db.get(models.Contact, cc.contact_id) if cc.contact_id else None
    email = (con.email if con and con.email else "").strip()
    return {
        "cc_id": cc.campaign_company_id,
        "campaign_id": camp.campaign_id,
        "contact_id": cc.contact_id,
        "company_id": cc.company_id,
        "email": email,
        "saluto": cc.saluto or "",
        "azienda": comp.legal_name if comp else "",
        "contesto": cc.contesto or "",
        "gancio": cc.gancio or "",
        "paragrafo_apertura": cc.paragrafo_apertura or build_apertura(cc.contesto or "", cc.gancio or ""),
    }


def _partiziona_pronte(db, camp):
    """Divide le righe Pronte fra (da_inviare, gia_contattati).

    Il controllo sui duplicati vive QUI, al momento dell'invio, e non solo
    all'import: cosi' vale per qualunque riga, anche per i file caricati
    prima che il controllo esistesse.
    """
    out, bloccati = [], []
    for cc in db.query(models.CampaignCompany).filter_by(campaign_id=camp.campaign_id).all():
        if cc.draft_status != "Pronta":
            continue
        comp = db.get(models.Company, cc.company_id)
        if comp and comp.suppression:
            continue
        con = db.get(models.Contact, cc.contact_id) if cc.contact_id else None
        if not con or not (con.email or "").strip() or con.opt_out:
            continue
        if db.query(models.OutreachMessage).filter_by(campaign_company_id=cc.campaign_company_id, message_no=1).first():
            continue
        precedente = learning.gia_contattato(db, con.email, camp.campaign_id)
        if precedente:
            quando, campagna = precedente
            bloccati.append({
                "cc_id": str(cc.campaign_company_id),
                "azienda": comp.legal_name if comp else "",
                "email": con.email,
                "quando": quando.strftime("%d/%m/%Y") if quando else "",
                "campagna": campagna or "",
            })
            continue
        out.append(_row_from_cc(db, cc, camp))
    return out, bloccati


def _ready_rows(db, camp):
    return _partiziona_pronte(db, camp)[0]


def _recall_rows(db, camp):
    out = []
    cutoff = datetime.now(timezone.utc) - timedelta(days=7)
    for cc in db.query(models.CampaignCompany).filter_by(campaign_id=camp.campaign_id).all():
        if cc.draft_status == "Esclusa":
            continue
        m1 = (db.query(models.OutreachMessage)
              .filter_by(campaign_company_id=cc.campaign_company_id, message_no=1, status="sent").first())
        if not m1 or not m1.sent_at or m1.sent_at > cutoff:
            continue
        if db.query(models.OutreachMessage).filter_by(campaign_company_id=cc.campaign_company_id, message_no=2).first():
            continue
        con = db.get(models.Contact, cc.contact_id) if cc.contact_id else None
        if not con or not (con.email or "").strip() or con.opt_out:
            continue
        if con.risposto_at:
            continue  # ha gia' risposto: il sollecito non si manda
        row = _row_from_cc(db, cc, camp)
        row["saluto_originale"] = saluto_orario(m1.sent_at)
        out.append(row)
    return out


@app.get("/campaign/{cid}/invio", response_class=HTMLResponse)
def invio_page(request: Request, cid: str, user=Depends(require_user)):
    db = SessionLocal()
    try:
        camp = db.get(models.Campaign, cid)
        if not camp:
            return RedirectResponse("/", status_code=303)
        rows, bloccati = _partiziona_pronte(db, camp)
    finally:
        db.close()
    return templates.TemplateResponse(request, "invio.html", {
        "camp": camp, "rows": rows, "tot": len(rows), "bloccati": bloccati,
        "state": mailer.STATE, "mode": "invio", "user": user,
        "auto_send": camp.auto_send, "auto_recall": camp.auto_recall, "inviate_oggi": mailer.inviate_oggi(),
        "nella_finestra": mailer.nella_finestra(), "daily_limit": DAILY_LIMIT,
        "win_start": SEND_WINDOW_START, "win_end": SEND_WINDOW_END,
    })


@app.get("/campaign/{cid}/recall", response_class=HTMLResponse)
def recall_page(request: Request, cid: str, user=Depends(require_user)):
    db = SessionLocal()
    try:
        camp = db.get(models.Campaign, cid)
        if not camp:
            return RedirectResponse("/", status_code=303)
        rows = _recall_rows(db, camp)
    finally:
        db.close()
    return templates.TemplateResponse(request, "invio.html", {
        "camp": camp, "rows": rows, "tot": len(rows), "bloccati": [],
        "state": mailer.STATE, "mode": "recall", "user": user,
        "auto_send": camp.auto_send, "auto_recall": camp.auto_recall, "inviate_oggi": mailer.inviate_oggi(),
        "nella_finestra": mailer.nella_finestra(), "daily_limit": DAILY_LIMIT,
        "win_start": SEND_WINDOW_START, "win_end": SEND_WINDOW_END,
    })


@app.post("/invio/test")
def invio_test(cid: str = Form(""), mode: str = Form("invio"), channel: str = Form("Gmail"),
               user=Depends(require_user)):
    db = SessionLocal()
    try:
        camp = db.get(models.Campaign, cid)
        rows = _recall_rows(db, camp) if mode == "recall" else _ready_rows(db, camp)
    finally:
        db.close()
    msg = mailer.do_test(rows, channel, 2 if mode == "recall" else 1)
    return RedirectResponse(f"/campaign/{cid}/{mode}?msg={msg}", status_code=303)


@app.post("/invio/start")
def invio_start(cid: str = Form(""), mode: str = Form("invio"), channel: str = Form("Gmail"),
                salva_copia: str = Form(""), user=Depends(require_user)):
    if mailer.STATE.get("running"):
        return RedirectResponse(f"/campaign/{cid}/{mode}?msg=Invio+gia+in+corso", status_code=303)
    db = SessionLocal()
    finestra = False
    try:
        camp = db.get(models.Campaign, cid)
        finestra = bool(camp.auto_send) if camp else False
        rows = _recall_rows(db, camp) if mode == "recall" else _ready_rows(db, camp)
    finally:
        db.close()
    if not rows:
        return RedirectResponse(f"/campaign/{cid}/{mode}?msg=Nessun+contatto+da+inviare", status_code=303)
    mailer.STOP.clear()
    import threading
    threading.Thread(target=mailer.run_send,
                     args=(rows, channel, salva_copia == "on", 2 if mode == "recall" else 1, finestra),
                     daemon=True).start()
    return RedirectResponse(f"/campaign/{cid}/{mode}", status_code=303)


@app.post("/campaign/{cid}/auto-send")
def auto_send_toggle(cid: str, attiva: str = Form("1"), mode: str = Form("invio"),
                     back: str = Form("/"), user=Depends(require_user)):
    db = SessionLocal()
    try:
        camp = db.get(models.Campaign, cid)
        if camp:
            if mode == "recall":
                camp.auto_recall = (attiva == "1")
            else:
                camp.auto_send = (attiva == "1")
            db.commit()
    finally:
        db.close()
    return RedirectResponse(back or f"/campaign/{cid}", status_code=303)


# ---------------------------------------------------------------------------
# Avvio programmato dall'esterno (Cloud Scheduler)
# ---------------------------------------------------------------------------
def _pulisci_token(valore):
    """Spazi, fine riga e BOM via dal token.

    Un segreto caricato su Secret Manager da PowerShell con `>` arriva in
    UTF-16 con BOM: senza questa pulizia il confronto fallirebbe sempre, e
    l'unico sintomo sarebbe un 403 senza spiegazione.
    """
    return (valore or "").strip().lstrip("\ufeff").strip()


@app.post("/scheduler/tick")
def scheduler_tick(request: Request, token: str = ""):
    """Un colpo di scheduler: al piu' una email, poi torna subito.

    Non usa il cookie di sessione, perche' chi chiama e' Cloud Scheduler e non
    un browser: si autentica con SCHEDULER_TOKEN, nell'intestazione
    X-Scheduler-Token oppure in ?token=. Senza token configurato, chiuso.
    """
    atteso = _pulisci_token(SCHEDULER_TOKEN)
    if not atteso:
        return JSONResponse({"errore": "SCHEDULER_TOKEN non configurato"}, status_code=503)
    fornito = _pulisci_token(request.headers.get("X-Scheduler-Token") or token)
    if not secrets.compare_digest(fornito, atteso):
        return JSONResponse({"errore": "token non valido"}, status_code=403)
    return JSONResponse(scheduler.tick())


@app.get("/scheduler/stato")
def scheduler_stato(user=Depends(require_user)):
    return JSONResponse(scheduler.stato_pubblico())


@app.post("/invio/stop")
def invio_stop(cid: str = Form(""), mode: str = Form("invio"), user=Depends(require_user)):
    mailer.STOP.set()
    return RedirectResponse(f"/campaign/{cid}/{mode}", status_code=303)


@app.get("/invio/status")
def invio_status(user=Depends(require_user)):
    with mailer.LOCK:
        return JSONResponse(dict(mailer.STATE))


# ---------------------------------------------------------------------------
# Risposte ricevute
# ---------------------------------------------------------------------------
@app.post("/risposte/scan")
def risposte_scan(giorni: str = Form(""), back: str = Form("/"),
                  user=Depends(require_user)):
    try:
        n = int(giorni) if giorni.strip() else None
    except ValueError:
        n = None
    # Parte in background: la pagina torna subito e segue l'avanzamento
    # da /risposte/status, senza rischiare il timeout della richiesta.
    avviata = replies.avvia_scan(n)
    sep = "&" if "?" in back else "?"
    if not avviata:
        return RedirectResponse(f"{back}{sep}msg={quote_plus('Scansione gia in corso')}",
                                status_code=303)
    return RedirectResponse(back or "/", status_code=303)


@app.get("/risposte/status")
def risposte_status(user=Depends(require_user)):
    with replies.LOCK:
        return JSONResponse(dict(replies.STATE))


if __name__ == "__main__":
    import uvicorn
    init_db()
    uvicorn.run("app.main:app", host=APP_HOST, port=APP_PORT, reload=False)
